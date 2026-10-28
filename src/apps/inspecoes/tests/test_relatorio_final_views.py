import io
import pytest
from datetime import date
from django.contrib.auth import get_user_model
from django.urls import reverse
from docx import Document

from apps.edificacoes.models import Edificacao, Profissional
from apps.inspecoes.models import Inspecao, InspecaoEspecialidade
from apps.inspecoes.views import _pode_gerar_relatorio_final

from django.core.files.uploadedfile import SimpleUploadedFile
from apps.inspecoes.models import Achado, Foto
from apps.inspecoes.views import _montar_snapshot_relatorio


def _prof(nome):
    return Profissional.objects.get_or_create(nome=nome, defaults={'crea': f'CREA-{nome}'})[0]


@pytest.fixture
def inspecao_com_profissionais(db):
    edif = Edificacao.objects.create(nome='Sede', descritivo='x')
    insp = Inspecao.objects.create(edificacao=edif)
    civil = InspecaoEspecialidade.objects.create(
        inspecao=insp, especialidade='civil',
        data_inspecao=date.today(), conclusao='ok',
    )
    civil.profissionais.set([_prof('Ana Civil')])
    eletrica = InspecaoEspecialidade.objects.create(
        inspecao=insp, especialidade='eletrica',
        data_inspecao=date.today(), conclusao='ok',
    )
    eletrica.profissionais.set([_prof('Beto Eletrica')])
    return insp


@pytest.mark.django_db
def test_profissional_de_qualquer_uma_das_especialidades_pode_gerar(inspecao_com_profissionais):
    U = get_user_model()
    civil = U.objects.create_user(username='ana', password='1', first_name='Ana', last_name='Civil')
    eletrica = U.objects.create_user(username='beto', password='1', first_name='Beto', last_name='Eletrica')
    estranho = U.objects.create_user(username='carlos', password='1', first_name='Carlos', last_name='Estranho')
    staff = U.objects.create_user(username='staff', password='1', is_staff=True)

    assert _pode_gerar_relatorio_final(civil, inspecao_com_profissionais)
    assert _pode_gerar_relatorio_final(eletrica, inspecao_com_profissionais)
    assert not _pode_gerar_relatorio_final(estranho, inspecao_com_profissionais)
    assert _pode_gerar_relatorio_final(staff, inspecao_com_profissionais)


@pytest.mark.django_db
def test_painel_mostra_pendencias_para_quem_tem_acesso(client, inspecao_com_profissionais):
    U = get_user_model()
    civil = U.objects.create_user(username='ana', password='1', first_name='Ana', last_name='Civil')
    client.force_login(civil)

    resp = client.get(reverse('inspecoes:relatorio_final_painel', kwargs={'pk': inspecao_com_profissionais.pk}))

    assert resp.status_code == 200
    assert 'Falta cadastrar' in resp.content.decode()  # falta Mecânica


@pytest.mark.django_db
def test_painel_nega_acesso_a_quem_nao_participou(client, inspecao_com_profissionais):
    U = get_user_model()
    estranho = U.objects.create_user(username='carlos', password='1', first_name='Carlos', last_name='Estranho')
    client.force_login(estranho)

    resp = client.get(reverse('inspecoes:relatorio_final_painel', kwargs={'pk': inspecao_com_profissionais.pk}), follow=True)

    assert resp.status_code == 200
    assert 'Acesso negado' in resp.content.decode()


@pytest.mark.django_db
def test_editar_descritivo_pelo_painel(client, inspecao_com_profissionais):
    U = get_user_model()
    civil = U.objects.create_user(username='ana', password='1', first_name='Ana', last_name='Civil')
    client.force_login(civil)

    resp = client.post(
        reverse('inspecoes:relatorio_final_editar_descritivo', kwargs={'pk': inspecao_com_profissionais.pk}),
        {'descritivo': 'Prédio de 4 pavimentos, estrutura em concreto armado.'},
    )

    assert resp.status_code == 302
    inspecao_com_profissionais.edificacao.refresh_from_db()
    assert inspecao_com_profissionais.edificacao.descritivo == 'Prédio de 4 pavimentos, estrutura em concreto armado.'


@pytest.mark.django_db
def test_snapshot_copia_fotos_e_resume_conformes(inspecao_com_profissionais):
    civil = inspecao_com_profissionais.especialidades.get(especialidade='civil')
    nc = Achado.objects.create(
        especialidade=civil, localizacao='L1', verificacao='Rachadura',
        grupo_tecnico='estrutura', requisito_afetado='seguranca_estrutural',
        gravidade=5, urgencia=5, tendencia=5, recomendacao='Reparar',
    )
    for i in range(3):  # 3 fotos — só as 2 mais antigas devem entrar
        Foto.objects.create(
            achado=nc, arquivo=SimpleUploadedFile(f'f{i}.jpg', f'conteudo{i}'.encode(), content_type='image/jpeg'),
            nome_original=f'f{i}.jpg',
        )
    conforme = Achado.objects.create(
        especialidade=civil, localizacao='L2', verificacao='Piso ok', grupo_tecnico='',
        requisito_afetado='', em_conformidade=True,
    )

    snapshot = _montar_snapshot_relatorio(inspecao_com_profissionais, numero_versao=1)

    civil_data = next(e for e in snapshot['especialidades'] if e['especialidade'] == 'civil')
    assert len(civil_data['achados_completos']) == 1
    assert len(civil_data['achados_completos'][0]['fotos']) == 2
    assert civil_data['achados_resumidos'] == [{'localizacao': 'L2', 'verificacao': 'Piso ok'}]

    # as fotos foram de fato copiadas para um caminho próprio do relatório
    from django.core.files.storage import default_storage
    for caminho in civil_data['achados_completos'][0]['fotos']:
        assert default_storage.exists(caminho)
        assert caminho.startswith(f'relatorios/{inspecao_com_profissionais.pk}/v1/')


@pytest.mark.django_db
def test_snapshot_inclui_dados_gerais_iguais_ao_painel_de_encerramento(inspecao_com_profissionais):
    """ADR da Seção 6 do design doc: o relatório reaproveita os mesmos
    números do painel de encerramento (análise geral) — não só a listagem
    itemizada de achados."""
    civil = inspecao_com_profissionais.especialidades.get(especialidade='civil')
    Achado.objects.create(
        especialidade=civil, localizacao='L1', verificacao='Rachadura',
        grupo_tecnico='estrutura', requisito_afetado='seguranca_estrutural',
        gravidade=5, urgencia=5, tendencia=5, prioridade_risco=1,
    )
    Achado.objects.create(
        especialidade=civil, localizacao='L2', verificacao='Piso ok', grupo_tecnico='',
        requisito_afetado='', em_conformidade=True,
    )

    snapshot = _montar_snapshot_relatorio(inspecao_com_profissionais, numero_versao=1)

    gerais = snapshot['dados_gerais']
    assert gerais['total_achados'] == 2
    assert gerais['total_nao_conformes'] == 1
    assert gerais['total_conformes'] == 1
    assert gerais['p1'] == 1
    civil_gerais = next(e for e in gerais['por_especialidade'] if e['especialidade_nome'] == 'Engenharia Civil')
    assert civil_gerais == {
        'especialidade_nome': 'Engenharia Civil', 'profissionais': 'Ana Civil',
        'data_inspecao': civil.data_inspecao.strftime('%d/%m/%Y'), 'status': 'Em andamento',
        'total': 2, 'total_nc': 1, 'p1': 1, 'p2': 0, 'p3': 0,
    }


@pytest.mark.django_db
def test_gerar_relatorio_final_cria_versao_1_e_bloqueia_sem_pendencias_resolvidas(client, inspecao_com_profissionais, settings):
    settings.GEMINI_API_KEY = ''  # sem chamada real ao Gemini neste teste
    U = get_user_model()
    civil = U.objects.create_user(username='ana', password='1', first_name='Ana', last_name='Civil')
    client.force_login(civil)

    # falta mecânica — deve ser bloqueado, nenhuma versão criada
    resp = client.post(reverse('inspecoes:relatorio_final_gerar', kwargs={'pk': inspecao_com_profissionais.pk}))
    assert resp.status_code == 302
    assert inspecao_com_profissionais.relatorios_finais.count() == 0

    mecanica_esp = InspecaoEspecialidade.objects.create(
        inspecao=inspecao_com_profissionais, especialidade='mecanica',
        data_inspecao=date.today(), conclusao='ok', status='finalizada',
    )
    mecanica_esp.profissionais.set([_prof('Carlos Mecanica')])
    inspecao_com_profissionais.especialidades.update(status='finalizada')

    resp = client.post(reverse('inspecoes:relatorio_final_gerar', kwargs={'pk': inspecao_com_profissionais.pk}))
    assert resp.status_code == 302
    assert inspecao_com_profissionais.relatorios_finais.count() == 1
    relatorio = inspecao_com_profissionais.relatorios_finais.first()
    assert relatorio.numero_versao == 1
    assert relatorio.gerado_por == civil
    assert relatorio.arquivo_docx.name
    assert relatorio.snapshot['edificacao_nome'] == 'Sede'

    # gerar de novo cria a versão 2, não sobrescreve a 1
    resp = client.post(reverse('inspecoes:relatorio_final_gerar', kwargs={'pk': inspecao_com_profissionais.pk}))
    assert inspecao_com_profissionais.relatorios_finais.count() == 2
    assert set(inspecao_com_profissionais.relatorios_finais.values_list('numero_versao', flat=True)) == {1, 2}

    from apps.inspecoes.models import LogAcesso
    assert LogAcesso.objects.filter(tipo='relatorio_final_gerado').count() == 2


@pytest.mark.django_db
def test_download_relatorio_final(client, inspecao_com_profissionais, settings):
    settings.GEMINI_API_KEY = ''  # sem chamada real ao Gemini neste teste
    U = get_user_model()
    civil = U.objects.create_user(username='ana', password='1', first_name='Ana', last_name='Civil')
    client.force_login(civil)
    mecanica_esp = InspecaoEspecialidade.objects.create(
        inspecao=inspecao_com_profissionais, especialidade='mecanica',
        data_inspecao=date.today(), conclusao='ok', status='finalizada',
    )
    mecanica_esp.profissionais.set([_prof('Carlos Mecanica')])
    inspecao_com_profissionais.especialidades.update(status='finalizada')
    client.post(reverse('inspecoes:relatorio_final_gerar', kwargs={'pk': inspecao_com_profissionais.pk}))
    relatorio = inspecao_com_profissionais.relatorios_finais.get()

    resp = client.get(reverse('inspecoes:relatorio_final_download', kwargs={'pk': inspecao_com_profissionais.pk, 'versao_pk': relatorio.pk}))

    assert resp.status_code == 200
    assert resp['Content-Type'] == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    # assinatura de arquivo .docx (zip) — PK\x03\x04
    assert resp.content.startswith(b'PK\x03\x04')


@pytest.mark.django_db
def test_reabrir_especialidade_avisa_quando_ja_ha_relatorio_gerado(client, inspecao_com_profissionais, settings):
    settings.GEMINI_API_KEY = ''  # sem chamada real ao Gemini neste teste
    U = get_user_model()
    civil_user = U.objects.create_user(username='ana', password='1', first_name='Ana', last_name='Civil')
    client.force_login(civil_user)
    mecanica_esp = InspecaoEspecialidade.objects.create(
        inspecao=inspecao_com_profissionais, especialidade='mecanica',
        data_inspecao=date.today(), conclusao='ok', status='finalizada',
    )
    mecanica_esp.profissionais.set([_prof('Carlos Mecanica')])
    inspecao_com_profissionais.especialidades.update(status='finalizada')
    client.post(reverse('inspecoes:relatorio_final_gerar', kwargs={'pk': inspecao_com_profissionais.pk}))

    civil = inspecao_com_profissionais.especialidades.get(especialidade='civil')
    resp = client.post(reverse('inspecoes:especialidade_reabrir', kwargs={'pk': civil.pk}), follow=True)

    assert resp.status_code == 200
    mensagens = [str(m) for m in resp.context['messages']]
    assert any('Relatório Final' in m and 'v1' in m for m in mensagens)


@pytest.mark.django_db
def test_editar_conclusao_de_especialidade_finalizada_avisa_quando_ja_ha_relatorio_gerado(client, inspecao_com_profissionais, settings):
    settings.GEMINI_API_KEY = ''  # sem chamada real ao Gemini neste teste
    U = get_user_model()
    civil_user = U.objects.create_user(username='ana', password='1', first_name='Ana', last_name='Civil')
    client.force_login(civil_user)
    mecanica_esp = InspecaoEspecialidade.objects.create(
        inspecao=inspecao_com_profissionais, especialidade='mecanica',
        data_inspecao=date.today(), conclusao='ok', status='finalizada',
    )
    mecanica_esp.profissionais.set([_prof('Carlos Mecanica')])
    inspecao_com_profissionais.especialidades.update(status='finalizada')
    client.post(reverse('inspecoes:relatorio_final_gerar', kwargs={'pk': inspecao_com_profissionais.pk}))

    civil = inspecao_com_profissionais.especialidades.get(especialidade='civil')
    resp = client.post(
        reverse('inspecoes:especialidade_conclusao', kwargs={'pk': civil.pk}),
        {'conclusao': 'Conclusão revisada após gerar o relatório.'},
        follow=True,
    )

    assert resp.status_code == 200
    civil.refresh_from_db()
    assert civil.conclusao == 'Conclusão revisada após gerar o relatório.'
    mensagens = [str(m) for m in resp.context['messages']]
    assert any('Relatório Final' in m and 'v1' in m for m in mensagens)


@pytest.mark.django_db
def test_especialidade_update_nao_aceita_mais_campo_conclusao(client, inspecao_com_profissionais):
    """A conclusão saiu de `EspecialidadeForm` (tela "Editar") para a tela
    dedicada "Conclusão" (`especialidade_conclusao`) — editar os campos
    cadastrais da especialidade não deve mais tocar na conclusão."""
    U = get_user_model()
    civil_user = U.objects.create_user(username='ana', password='1', first_name='Ana', last_name='Civil')
    client.force_login(civil_user)

    civil = inspecao_com_profissionais.especialidades.get(especialidade='civil')
    civil.conclusao = 'Conclusão original.'
    civil.save(update_fields=['conclusao'])

    prof_pk = _prof('Ana Civil').pk
    resp = client.post(
        reverse('inspecoes:especialidade_update', kwargs={'pk': civil.pk}),
        {
            'especialidade': 'civil',
            'data_inspecao': civil.data_inspecao.isoformat(),
            'conclusao': 'Tentativa de mudar a conclusão por aqui.',
            'profissionais': [str(prof_pk)],
        },
        follow=True,
    )

    assert resp.status_code == 200
    civil.refresh_from_db()
    assert civil.conclusao == 'Conclusão original.'


@pytest.mark.django_db
def test_relatorio_final_gerar_produz_docx(client, django_user_model, settings, inspecao_com_profissionais):
    """Teste de integração da Task 11: a view `relatorio_final_gerar`
    monta o contexto completo (dados determinísticos + textos do Gemini,
    aqui vazios por falta de API key) e renderiza o template .docx real,
    sem deixar nenhum placeholder `{{ }}` sem substituir."""
    settings.GEMINI_API_KEY = ''  # sem chamada real ao Gemini neste teste
    civil_user = django_user_model.objects.create_user(
        username='ana', password='1', first_name='Ana', last_name='Civil',
    )
    client.force_login(civil_user)

    mecanica_esp = InspecaoEspecialidade.objects.create(
        inspecao=inspecao_com_profissionais, especialidade='mecanica',
        data_inspecao=date.today(), conclusao='ok', status='finalizada',
    )
    mecanica_esp.profissionais.set([_prof('Carlos Mecanica')])
    inspecao_com_profissionais.especialidades.update(status='finalizada')

    civil = inspecao_com_profissionais.especialidades.get(especialidade='civil')
    nc = Achado.objects.create(
        especialidade=civil, localizacao='L1', verificacao='Rachadura na parede',
        grupo_tecnico='estrutura', requisito_afetado='seguranca_estrutural',
        gravidade=5, urgencia=5, tendencia=5, recomendacao='Reparar a fissura.',
        prioridade_risco=1, prazo_meses=1,
    )
    Foto.objects.create(
        achado=nc, arquivo=SimpleUploadedFile('f1.jpg', b'conteudo1', content_type='image/jpeg'),
        nome_original='f1.jpg',
    )
    Achado.objects.create(
        especialidade=civil, localizacao='L2', verificacao='Piso ok', grupo_tecnico='acabamento',
        requisito_afetado='estetica', em_conformidade=True,
    )

    inspecao = inspecao_com_profissionais
    resp = client.post(reverse('inspecoes:relatorio_final_gerar', kwargs={'pk': inspecao.pk}))

    assert resp.status_code == 302
    relatorio = inspecao.relatorios_finais.first()
    assert relatorio is not None
    assert relatorio.arquivo_docx.name.endswith('.docx')

    doc = Document(io.BytesIO(relatorio.arquivo_docx.read()))
    texto = '\n'.join(p.text for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                texto += '\n' + '\n'.join(p.text for p in cell.paragraphs)
    assert '{{' not in texto
    assert '{%' not in texto


@pytest.mark.django_db
def test_falha_na_geracao_do_docx_limpa_fotos_ja_copiadas(client, inspecao_com_profissionais, settings):
    """Cobre o achado da revisão de qualidade: `_montar_snapshot_relatorio`
    (chamado antes de `gerar_docx_relatorio`) já copiou as fotos para
    relatorios/<pk>/v<N>/... antes de sabermos se o .docx vai ser gerado com
    sucesso. Se `gerar_docx_relatorio` lançar uma exceção, essas cópias não
    podem ficar órfãs — nenhuma RelatorioFinalInspecao vai apontar pra elas.

    Diferente da antiga versão em PDF (que usava um retorno sentinela
    `None` engolido pela view), `gerar_docx_relatorio` não tem esse padrão —
    uma falha real é uma exceção, que deve propagar (não ser convertida
    silenciosamente numa mensagem de erro + redirect). Por isso o teste
    força a falha via `unittest.mock.patch` com `side_effect` e espera que
    a própria chamada a `client.post` levante a exceção (comportamento
    padrão do test client do Django: exceções não tratadas pela view
    propagam para o teste).

    Captura os caminhos reais via um espião em `_montar_snapshot_relatorio`
    em vez de prever o nome do arquivo: o storage de mídia deste projeto não
    é isolado por teste, então um arquivo de uma execução anterior pode já
    ocupar o nome "óbvio" e forçar o Django a usar um sufixo diferente.
    """
    from unittest.mock import patch
    from django.core.files.storage import default_storage
    import apps.inspecoes.views as views_module

    settings.GEMINI_API_KEY = ''  # sem chamada real ao Gemini neste teste

    U = get_user_model()
    civil_user = U.objects.create_user(username='ana3', password='1', first_name='Ana', last_name='Civil')
    client.force_login(civil_user)

    mecanica_esp = InspecaoEspecialidade.objects.create(
        inspecao=inspecao_com_profissionais, especialidade='mecanica',
        data_inspecao=date.today(), conclusao='ok', status='finalizada',
    )
    mecanica_esp.profissionais.set([_prof('Carlos Mecanica')])
    inspecao_com_profissionais.especialidades.update(status='finalizada')
    civil = inspecao_com_profissionais.especialidades.get(especialidade='civil')
    achado = Achado.objects.create(
        especialidade=civil, localizacao='L1', verificacao='Rachadura',
        grupo_tecnico='estrutura', requisito_afetado='seguranca_estrutural',
        gravidade=5, urgencia=5, tendencia=5,
    )
    Foto.objects.create(
        achado=achado, arquivo=SimpleUploadedFile('f1.jpg', b'conteudo1', content_type='image/jpeg'),
        nome_original='f1.jpg',
    )
    Foto.objects.create(
        achado=achado, arquivo=SimpleUploadedFile('f2.jpg', b'conteudo2', content_type='image/jpeg'),
        nome_original='f2.jpg',
    )

    capturado = {}
    original_montar_snapshot = views_module._montar_snapshot_relatorio

    def _snapshot_espiao(inspecao, numero_versao):
        resultado = original_montar_snapshot(inspecao, numero_versao)
        capturado['snapshot'] = resultado
        return resultado

    with patch('apps.inspecoes.views._montar_snapshot_relatorio', side_effect=_snapshot_espiao), \
         patch('apps.inspecoes.views.gerar_docx_relatorio', side_effect=RuntimeError('falha simulada na geração do docx')):
        with pytest.raises(RuntimeError, match='falha simulada na geração do docx'):
            client.post(reverse('inspecoes:relatorio_final_gerar', kwargs={'pk': inspecao_com_profissionais.pk}))

    assert inspecao_com_profissionais.relatorios_finais.count() == 0

    snapshot = capturado['snapshot']
    civil_data = next(e for e in snapshot['especialidades'] if e['especialidade'] == 'civil')
    caminhos_copiados = civil_data['achados_completos'][0]['fotos']
    assert len(caminhos_copiados) == 2  # sanity: as fotos foram de fato copiadas antes da falha
    for caminho in caminhos_copiados:
        assert not default_storage.exists(caminho)
